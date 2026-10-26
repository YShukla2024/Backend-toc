import os
import re
from pathlib import Path
from io import BytesIO
import chromadb
from docx import Document
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer


class VectorService:

    _model = None

    def __init__(self):

        BASE_DIR = Path(__file__).resolve().parent.parent
        storage_path = BASE_DIR / "data" / "embeddings"
        storage_path.mkdir(parents=True, exist_ok=True)

        print("✅ DB PATH:", storage_path.resolve())

        self.client = chromadb.PersistentClient(path=str(storage_path))

        # 🔥 Keep light logs (avoid heavy startup)
        print("Chroma initialized")

        self.collection = self.client.get_or_create_collection(
            name="toc_repository"
        )

        print("📦 VECTOR COUNT:", self.collection.count())

    # --------------------------------
    # ✅ LAZY MODEL LOADING (FIX)
    # --------------------------------
    @classmethod
    def get_model(cls):
        if cls._model is None:
            print("🔄 Loading embedding model...")
            cls._model = SentenceTransformer(
                "sentence-transformers/all-MiniLM-L6-v2"
            )
        return cls._model

    # ------------------------------
    # Extract text from DOCX
    # ------------------------------
    def extract_text_from_docx(self, file_path):

        doc = Document(file_path)
        full_text = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                full_text.append(text)

        return "\n".join(full_text)

    # ------------------------------
    # Chunk text
    # ------------------------------
    def chunk_text(self, text, chunk_size=800, overlap=100):

        chunks = []
        start = 0

        while start < len(text):
            end = start + chunk_size
            chunks.append(text[start:end])
            start += chunk_size - overlap

        return chunks

    # ------------------------------
    # Smart chunk
    # ------------------------------
    def smart_chunk(self, text):

        chunks = []

        day_sections = re.split(r'(Day\s*\d+)', text, flags=re.IGNORECASE)

        if len(day_sections) > 1:
            for i in range(1, len(day_sections), 2):
                title = day_sections[i]
                content = day_sections[i+1] if i+1 < len(day_sections) else ""
                chunk = (title + content).strip()

                if len(chunk) > 50:
                    chunks.append(chunk)

            return chunks

        module_sections = re.split(r'(Module[:\s].*)', text, flags=re.IGNORECASE)

        if len(module_sections) > 1:
            for i in range(1, len(module_sections), 2):
                title = module_sections[i]
                content = module_sections[i+1] if i+1 < len(module_sections) else ""
                chunk = (title + content).strip()

                if len(chunk) > 50:
                    chunks.append(chunk)

            return chunks

        paragraphs = text.split("\n\n")
        return [p.strip() for p in paragraphs if len(p.strip()) > 50]

    # ------------------------------
    # Index DOCX file
    # ------------------------------
    def index_document(self, file_path):

        text = self.extract_text_from_docx(file_path)

        if not text:
            print("No text found")
            return

        chunks = self.smart_chunk(text)

        model = VectorService.get_model()

        embeddings = model.encode(
            chunks,
            batch_size=32,
            show_progress_bar=True
        ).tolist()

        file_name = os.path.basename(file_path)

        ids = [f"{file_name}_{i}" for i in range(len(chunks))]
        metadatas = [{"source": file_name, "chunk": i} for i in range(len(chunks))]

        self.collection.add(
            documents=chunks,
            embeddings=embeddings,
            ids=ids,
            metadatas=metadatas
        )

        print(f"Indexed {len(chunks)} chunks")

    # ------------------------------
    # Extract DOCX bytes
    # ------------------------------
    def extract_text_from_docx_bytes(self, file_bytes):

        doc = Document(BytesIO(file_bytes))
        full_text = []

        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        full_text.append(text)

        return "\n".join(full_text)

    # ------------------------------
    # Extract PDF bytes
    # ------------------------------
    def extract_text_from_pdf_bytes(self, file_bytes):

        reader = PdfReader(BytesIO(file_bytes))
        full_text = []

        for page in reader.pages:
            text = page.extract_text()
            if text:
                full_text.append(text.strip())

        return "\n".join(full_text)

    # ------------------------------
    # Index from bytes
    # ------------------------------
    def index_document_from_bytes(self, file_bytes, file_name):

        existing_sources = set(
            m["source"] for m in self.collection.get()["metadatas"] if m
        )

        if file_name in existing_sources:
            print("Already indexed")
            return

        text = ""

        if file_name.lower().endswith(".docx"):
            text = self.extract_text_from_docx_bytes(file_bytes)

        elif file_name.lower().endswith(".pdf"):
            text = self.extract_text_from_pdf_bytes(file_bytes)

        else:
            return

        if not text:
            return

        chunks = self.smart_chunk(text)

        model = VectorService.get_model()

        embeddings = model.encode(
            chunks,
            batch_size=32,
            show_progress_bar=False
        ).tolist()

        ids = [f"{file_name}_{i}" for i in range(len(chunks))]
        metadatas = [{"source": file_name, "chunk": i} for i in range(len(chunks))]

        self.collection.add(
            documents=chunks,
            embeddings=embeddings,
            ids=ids,
            metadatas=metadatas
        )

        print(f"Indexed {len(chunks)} chunks")

    # ------------------------------
    # Search
    # ------------------------------
    def search(self, query, top_k=5):

        print("VECTOR COUNT:", self.collection.count())
        print("SEARCH QUERY:", query)

        model = VectorService.get_model()

        query_embedding = model.encode(query).tolist()

        results = self.collection.query(
            query_embeddings=[query_embedding],
            n_results=top_k
        )

        documents = results.get("documents", [[]])[0]
        metadatas = results.get("metadatas", [[]])[0]
        distances = results.get("distances", [[]])[0]

        results_list = []

        for doc, meta, dist in zip(documents, metadatas, distances):
            results_list.append({
                "text": doc,
                "source": meta.get("source", "unknown") if meta else "unknown",
                "chunk": meta.get("chunk", -1) if meta else -1,
                "distance": dist
            })

        return results_list